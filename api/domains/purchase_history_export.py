from datetime import datetime, time, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


HISTORY_TIMEZONE = ZoneInfo("Europe/Moscow")
HEADERS = [
    "ID покупки", "Источник", "ID услуги", "Сервис", "ID номинала", "Номинал",
    "Цена, ₽", "Статус", "Дата (МСК)", "Сделка", "Заказ", "Покупатель",
    "Регион", "Оператор", "Гифт-код CRM",
]
STATES = {"paid": "Оплачено", "succeeded": "Выполнено", "processing": "В обработке",
          "requires_attention": "Нужна проверка", "failed": "Не выполнено"}


def history_date(value):
    # Сохраняем дату числом Excel в МСК, чтобы её можно было сортировать и фильтровать.
    if not value:
        return None
    parsed = value if isinstance(value, datetime) else datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=HISTORY_TIMEZONE)
    return parsed.astimezone(HISTORY_TIMEZONE).replace(tzinfo=None)


def build_purchase_history_xlsx(*, load_crm_rows, supplier_hub_client, services, date_from, date_to):
    # Выгружаем успешные покупки отдельно и на общий лист, читая каждый источник только один раз.
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title in ("CRM", "Селлер", "Все покупки"):
        workbook.create_sheet(title).append(HEADERS)
    combined_sheet = workbook["Все покупки"]
    combined_row_number = 1
    labels = {}
    nominals = {}
    for service in services:
        service_id = int(service.get("service_id") or 0)
        labels[service_id] = service.get("title") or str(service_id)
        for field in service.get("fields") or []:
            if field.get("name") == "nominal":
                for nominal in field.get("value_list") or []:
                    nominals[(service_id, str(nominal.get("id")))] = nominal.get("title")

    # В CRM архив уже ограничен paid, а у селлера успешной покупке соответствует succeeded.
    query = {"state": "succeeded", "sort_by": "created_at", "sort_direction": "asc"}
    if date_from:
        query["created_from"] = datetime.combine(date_from, time.min, HISTORY_TIMEZONE).isoformat()
    if date_to:
        query["created_to"] = datetime.combine(date_to + timedelta(days=1), time.min, HISTORY_TIMEZONE).isoformat()

    for source, items in (("CRM", load_crm_rows()), ("Селлер", iter_supplier_history(supplier_hub_client, query))):
        sheet = workbook[source]
        # Номер строки считаем сами: max_row при каждом обращении обходит все накопленные ячейки.
        for row_number, item in enumerate(items, start=2):
            service_id = int(item.get("service_id") or 0)
            nominal_id = str(item.get("nominal") or item.get("nominal_id") or "")
            state = item.get("state") or ("paid" if source == "CRM" else "")
            amount = item.get("price") if source == "CRM" else item.get("amount")
            if amount is None:
                amount = item.get("max_amount")
            values = [
                str(item.get("agent_transaction_id") or item.get("id") or ""),
                "CRM" if source == "CRM" else str(item.get("consumer_id") or "seller"),
                service_id, item.get("service_title") or labels.get(service_id, str(service_id)),
                nominal_id, item.get("nominal_title") or nominals.get((service_id, nominal_id)) or nominal_id,
                float(amount or 0),
                STATES.get(state, state), history_date(item.get("created_at")),
                item.get("deal_id"), item.get("order_number", ""), item.get("customer_nickname", ""),
                item.get("region_code", ""), item.get("created_by", ""),
                item.get("gift_code", "") if source == "CRM" else "",
            ]
            combined_row_number += 1
            # Одну и ту же строку записываем на оба листа с одинаковыми типами ячеек и защитой от формул.
            for target, target_row in ((sheet, row_number), (combined_sheet, combined_row_number)):
                target.append(values)
                for col in range(1, len(HEADERS) + 1):
                    cell = target.cell(target_row, col)
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                target.cell(target_row, 7).number_format = '#,##0.00'
                target.cell(target_row, 9).number_format = 'dd.mm.yyyy hh:mm:ss'
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243746")
        for index, width in enumerate([38, 18, 14, 32, 18, 26, 16, 22, 22, 14, 22, 24, 14, 22, 32], 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def iter_supplier_history(client, query):
    # Читаем успешные покупки селлера всеми страницами и не выдаём незаметно обрезанный файл.
    offset = 0
    while True:
        payload = client.list_transactions({**query, "limit": 100, "offset": offset})
        items = payload.get("items") or []
        total = int(payload["total"])
        if not items and offset < total:
            raise HTTPException(502, "История покупок загружена не полностью. Повторите выгрузку")
        yield from items
        offset += len(items)
        if offset >= total:
            return
