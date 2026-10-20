"""Отчёт о проданных ваучерах: Seller → сверка CRM → Excel без открытых кодов."""
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
import hashlib
from io import BytesIO
import json
import os
import re
import urllib.error
import urllib.parse
import urllib.request
from zoneinfo import ZoneInfo

from fastapi import Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

MSK = ZoneInfo('Europe/Moscow')
HEADERS = ['Номер заказа', 'Название сервиса', 'Номинал', 'Закупочная цена, ₽', 'Дата продажи', 'Комментарий']
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def report_window(first, last):
    # Проверяем календарные границы до доступа к источникам и включаем последний день.
    now = datetime.now(timezone.utc)
    if first > last or (last-first).days >= 366 or last > now.astimezone(MSK).date():
        raise HTTPException(422, 'Укажите период до сегодняшнего дня, не более 366 дней')
    return datetime.combine(first, time.min, MSK), min(datetime.combine(last+timedelta(days=1), time.min, MSK), now)


class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        # Не пересылаем машинный секрет на адрес из переадресации.
        return None


def read_seller(first, last):
    # CRM не получает доступ к БД Seller: используется отдельный ограниченный HTTPS-контракт.
    base = os.getenv('SELLER_REPORT_URL', '').rstrip('/')
    token = os.getenv('SELLER_CRM_REPORT_TOKEN', '')
    url = urllib.parse.urlsplit(base)
    if (len(token) < 32 or not url.hostname or url.username or url.password or url.query or url.fragment
            or url.path not in ('', '/', '/api') or (url.scheme != 'https' and not
            (url.scheme == 'http' and url.hostname in ('localhost', '127.0.0.1')))):
        raise HTTPException(503, 'Отчёт не настроен: требуется подключение к Seller')
    query = urllib.parse.urlencode({'date_from': first.isoformat(), 'date_to': last.isoformat()})
    req = urllib.request.Request(base+'/internal/reports/voucher-sales?'+query,
                                 headers={'X-Voucher-Report-Token': token})
    try:
        with urllib.request.build_opener(NoRedirect).open(req, timeout=60) as response:
            raw = response.read(20*1024*1024+1)
        if len(raw) > 20*1024*1024:
            raise ValueError('size')
        payload = json.loads(raw)
        if (payload.get('version') != 1 or not isinstance(payload.get('rows'), list)
                or payload.get('total') != len(payload['rows']) or len(payload['rows']) > 20000
                or payload.get('date_from') != first.isoformat() or payload.get('date_to') != last.isoformat()):
            raise ValueError('contract')
        return payload
    except (urllib.error.URLError, OSError, ValueError, TypeError) as exc:
        raise HTTPException(502, 'Не удалось получить полный отчёт Seller. Попробуйте позже или уменьшите период') from exc


def normalized_hash(code):
    # Убираем только разделители печатного кода, не выводя сам код из процесса.
    return hashlib.sha256(re.sub(r'[\s-]', '', code).upper().encode()).hexdigest()


def load_catalog(cur):
    # Читаем текущие и исторические подписи, поэтому исчезновение номинала не меняет прошлую продажу.
    cur.execute('''SELECT DISTINCT ON (service_id,nominal_id) service_id,nominal_id,service_title,nominal_title
        FROM app.supplier_catalog_labels WHERE success=true
        ORDER BY service_id,nominal_id,calculated_at DESC,id DESC''')
    catalog = {(str(sid),str(nid)): (title,nominal) for sid,nid,title,nominal in cur.fetchall()}
    cur.execute("SELECT service_id,nominal_id,service_title,nominal_title,service_type FROM app.supplier_catalog_current WHERE supplier_code='interhub'")
    types = {}
    for sid,nid,title,nominal,kind in cur.fetchall():
        catalog[(str(sid),str(nid))] = (title,nominal)
        types[str(sid)] = kind
    return catalog, types


def match_purchases(cur, rows, cutoff):
    # Архив читаем пачками один раз: закупка из пула могла состояться задолго до продажи.
    wanted = defaultdict(list)
    for row in rows:
        f = row.get('fingerprints')
        if f:
            for value in {f['whole'], f['normalized'], *f['tokens']}:
                wanted[value].append(row['key_id'])
    matches = defaultdict(dict)
    if not wanted:
        return matches
    with cur.connection.cursor(name='voucher_report_archive') as archive:
        archive.execute('''SELECT agent_transaction_id,service_id,request_params->>'nominal',amount,
            created_at,gift_code FROM app.interhub_transactions
            WHERE state='paid' AND created_at<%s AND COALESCE(gift_code,'')<>''
            ORDER BY created_at,agent_transaction_id''', (cutoff,))
        scanned = 0
        while batch := archive.fetchmany(500):
            scanned += len(batch)
            if scanned > 500000:
                raise HTTPException(422, 'Архив закупок слишком большой для синхронного отчёта')
            for tx,sid,nid,amount,created,code in batch:
                raw_hash = hashlib.sha256(code.strip().encode()).hexdigest()
                norm_hash = normalized_hash(code)
                for key in set(wanted.get(raw_hash, []) + wanted.get(norm_hash, [])):
                    matches[key]['interhub:'+tx] = dict(id=tx, service_id=str(sid), nominal_id=str(nid),
                        amount=amount, created_at=created, hashes={raw_hash,norm_hash}, provider='interhub')
    # Старый локальный AirPay хранит готовый SHA256. Его коды расшифровывать не требуется.
    cur.execute('''SELECT agent_transaction_id,amount,currency,created_at,result_hash
        FROM app.airpay_transactions WHERE state='paid' AND created_at<%s AND result_hash=ANY(%s)''',
        (cutoff, list(wanted)))
    for tx,amount,currency,created,hashed in cur.fetchall():
        if currency != 'RUB':
            continue
        for key in wanted[hashed]:
            matches[key]['airpay:'+tx] = dict(id=tx, amount=amount, created_at=created,
                hashes={hashed}, provider='airpay')
    return matches


def parse_stamp(value):
    # Отказываемся от дат без зоны: иначе календарная граница отчёта была бы неоднозначной.
    parsed = value if isinstance(value, datetime) else datetime.fromisoformat(value.replace('Z','+00:00'))
    if parsed.tzinfo is None:
        raise ValueError('timezone required')
    return parsed


def nominal_from_title(title):
    # Сохранённое название заказа помогает старым номиналам, которых уже нет в каталоге.
    found = re.search(r'\b(\d+(?:[.,]\d+)?)\s*(USD|EUR|TRY|AED|GBP|PLN|RUB)\b', title, re.I)
    if found:
        return found[1].replace(',', '.')+' '+found[2].upper()
    return ''


def assemble_rows(sales, matches, catalog, types, start, end):
    # Подтверждаем только полное и однозначное покрытие кодов; время и номинал сами по себе не доказывают закупку.
    result = []
    used = Counter(tx for values in matches.values() for tx in values)
    seen = set()
    for row in sales:
        key = row['key_id']
        if key in seen or not start <= parse_stamp(row['delivered_at']) < end:
            raise ValueError('invalid sales rows')
        seen.add(key)
        if row['delivery_source']=='supplier' and row.get('provider_code') not in (None,'interhub'):
            continue
        sid,nid = str(row.get('service_id') or ''),str(row.get('nominal_id') or '')
        if types.get(sid) not in (None,'VOUCHER'):
            continue
        service,nominal = catalog.get((sid,nid), ('',''))
        service = re.sub(r'^po_', '', service) or row['title']
        nominal = re.sub(r'^([A-Z]{3}) (.+)$', r'\2 \1', nominal) or nominal_from_title(row['title']) or 'Не определён'
        # В ручных комплектах номинал продажи берётся из заказа, не из одной закупки комплекта.
        if row['delivery_source'] in ('manual','pool'):
            nominal = nominal_from_title(row['title']) or nominal
        amount = row.get('amount') if row.get('attempt_id') and row.get('provider_code')=='interhub' else None
        kind = 'supplier'
        comment = detail = ''
        if row['delivery_source'] in ('manual','pool'):
            kind = row['delivery_source']
            f = row['fingerprints']
            found = {tx:m for tx,m in matches.get(key, {}).items()
                     if parse_stamp(m['created_at']) <= parse_stamp(row['delivered_at'])}
            whole = [(tx,m) for tx,m in found.items() if {f['whole'],f['normalized']} & m['hashes']]
            covered = set().union(*(m['hashes'] for m in found.values())) if found else set()
            complete = (len(whole)==1 and len(found)==1) or (not whole and bool(f['tokens']) and set(f['tokens']) <= covered)
            unique = all(used[tx]==1 for tx in found)
            known = all(m['amount'] is not None and Decimal(str(m['amount'])).is_finite() for m in found.values())
            # Один токен не может иметь две разные закупки с одинаковым кодом.
            unambiguous = all(sum(h in m['hashes'] for m in found.values())==1 for h in f['tokens'] if h in covered)
            subtotal = sum((Decimal(str(m['amount'])) for m in found.values() if m['amount'] is not None), Decimal(0))
            amount = subtotal if complete and unique and known and unambiguous else None
            prefix = 'Из пула' if kind=='pool' else 'Ручная выдача'
            if amount is not None:
                comment = prefix+' — закупка подтверждена по коду.'
            elif found:
                comment = prefix+' — полная закупочная цена не подтверждена.'
            else:
                comment = prefix+' — закупка не найдена.'
            detail = comment
            if found:
                detail += ' Покупки: '+', '.join(tx for tx in found)+'.'
                if amount is None:
                    detail += f' Найденная часть: {subtotal:.2f} ₽. В полную стоимость не включена.'
        if amount is None:
            kind = 'unconfirmed'
            comment = comment or 'Закупочная цена не подтверждена.'
        else:
            amount = Decimal(str(amount))
            if not amount.is_finite() or amount < 0:
                raise ValueError('invalid amount')
        result.append(dict(connection_id=row['connection_id'],shop=row['shop'],order=str(row['order_id']),
            service=service,nominal=nominal,amount=amount,date=parse_stamp(row['delivered_at']).astimezone(MSK).replace(tzinfo=None),
            kind=kind,comment=comment,detail=detail))
    return result


def build_xlsx(rows):
    # В файле только название магазина и таблица; пустых магазинов, дат периода и счётчиков нет.
    if not rows:
        raise HTTPException(404, 'За выбранный период продаж ваучеров нет')
    workbook = Workbook(); workbook.remove(workbook.active)
    groups = defaultdict(list)
    for row in rows:
        groups[row['connection_id']].append(row)
    names = set()
    for cid,items in groups.items():
        title = re.sub(r'[\\/*?:\[\]]', '_', items[0]['shop']).strip("'")[:31] or 'Магазин'
        base = title; index = 1
        while title.lower() in names:
            index += 1; suffix = f' ({index})'; title = base[:31-len(suffix)]+suffix
        names.add(title.lower()); sheet = workbook.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        sheet.append([items[0]['shop']]); sheet['A1'].font = Font(name='Arial',size=15,bold=True)
        sheet.append(HEADERS)
        for cell in sheet[2]:
            cell.fill = PatternFill('solid',fgColor='243B53'); cell.font = Font(name='Arial',bold=True,color='FFFFFF')
        for item in items:
            sheet.append([item['order'],item['service'],item['nominal'],item['amount'] if item['amount'] is not None else 'Не подтверждена',item['date'],item['comment']])
            rownum = sheet.max_row
            color = {'manual':'FFF1C9','pool':'ECE3F8','unconfirmed':'FBE4DF'}.get(item['kind'])
            for cell in sheet[rownum]:
                cell.font = Font(name='Arial',size=11)
                cell.alignment = Alignment(vertical='center')
                if color: cell.fill = PatternFill('solid',fgColor=color)
            sheet.cell(rownum,4).number_format = '#,##0.00'
            sheet.cell(rownum,5).number_format = 'dd.mm.yyyy hh:mm'
            if item['detail']:sheet.cell(rownum,6).comment = Comment(item['detail'],'CRM')
        # Названия и номера могут начинаться с «=»; они всегда остаются текстом Excel.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
        for col,width in enumerate([24,48,20,24,24,65],1):
            sheet.column_dimensions[get_column_letter(col)].width=width
        sheet.freeze_panes='A3';sheet.auto_filter.ref=f'A2:F{sheet.max_row}'
    output=BytesIO();workbook.save(output);return output.getvalue()


def mount_voucher_sales_report(app, *, db, dsn, require_role):
    # Существующее право владельца CRM защищает выгрузку закупочных цен и сверку материалов.
    @app.get('/integrations/interhub/voucher-sales/export')
    def export(date_from: date = Query(), date_to: date = Query(), _user=Depends(require_role('owner'))):
        start,end = report_window(date_from,date_to)
        try:
            with db.connect(dsn) as conn, conn.cursor() as cur:
                # Пул делает SELECT 1 при выдаче соединения; начинаем свою read-only транзакцию после него.
                conn.rollback()
                cur.execute('SET TRANSACTION READ ONLY')
                cur.execute("SET LOCAL statement_timeout='45s'")
                cur.execute("SELECT pg_try_advisory_xact_lock(hashtext('crm-voucher-sales-report'))")
                if not cur.fetchone()[0]:
                    raise HTTPException(409,'Отчёт уже формируется. Повторите выгрузку после завершения')
                payload=read_seller(date_from,date_to)
                cutoff=parse_stamp(payload['cutoff'])
                if not start < cutoff <= end+timedelta(minutes=1):
                    raise ValueError('cutoff')
                catalog,types=load_catalog(cur)
                matches=match_purchases(cur,payload['rows'],cutoff)
                rows=assemble_rows(payload['rows'],matches,catalog,types,start,cutoff)
                content=build_xlsx(rows)
        except HTTPException:
            raise
        except Exception as exc:
            # Ошибка источника не превращается в частичный файл или ответ с кодами и DSN.
            raise HTTPException(502,'Не удалось сформировать полный отчёт. Повторите позже') from exc
        return Response(content,media_type=XLSX,headers={
            'Content-Disposition':f'attachment; filename="voucher-sales-{date_from}-{date_to}.xlsx"',
            'Cache-Control':'no-store'})
