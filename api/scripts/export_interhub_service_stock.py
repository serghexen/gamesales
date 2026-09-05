"""Выгружает остатки InterHub по активным услугам через service/detail в Excel."""

from __future__ import annotations

import argparse
import json
import os
import re
import socket
import ssl
import sys
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Callable
from zoneinfo import ZoneInfo

from dotenv import load_dotenv
from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def resolve_runtime_paths(script_path: Path, current_dir: Path) -> tuple[Path, Path]:
    # Находит API-модули и в репозитории, и при временном запуске скрипта внутри контейнера.
    resolved_script = script_path.resolve()
    for parent in resolved_script.parents:
        if (parent / "api" / "domains").is_dir():
            return parent, parent / "api"
    runtime_dir = current_dir.resolve()
    return runtime_dir, runtime_dir


ROOT_DIR, API_DIR = resolve_runtime_paths(Path(__file__), Path.cwd())
if str(API_DIR) not in sys.path:
    sys.path.insert(0, str(API_DIR))

from domains.interhub_service import build_interhub_service  # noqa: E402


STOCK_HEADERS = [
    "ID услуги",
    "Услуга",
    "Категория",
    "Тип",
    "ID нашего номинала",
    "Наш номинал",
    "Название в service/detail",
    "Остаток",
    "Результат сопоставления",
]


def parse_bool(value: Any, default: bool = False) -> bool:
    # Приводит строковую настройку окружения к булеву значению без неявных вариантов.
    if value is None:
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


def normalize_nominal_name(value: Any) -> str:
    # Нормализует регистр, Unicode и пробелы, сохраняя значимые цифры и символы номинала.
    normalized = unicodedata.normalize("NFKC", str(value or "")).replace("\u00a0", " ").casefold()
    return re.sub(r"\s+", " ", normalized).strip()


def canonical_nominal_name(value: Any) -> str:
    # Убирает только оформительские различия: разделители тысяч, нулевые копейки и пунктуацию.
    normalized = normalize_nominal_name(value)
    normalized = re.sub(r"(?<=\d)[.,](?=\d{3}(?:\D|$))", "", normalized)
    normalized = re.sub(r"(?<=\d)[.,]00\b", "", normalized)
    normalized = re.sub(r"[^\w]+", " ", normalized, flags=re.UNICODE)
    return re.sub(r"\s+", " ", normalized).strip()


def nominal_number_signature(value: Any) -> tuple[int, ...]:
    # Извлекает числа из имени для осторожного резервного сопоставления внутри одной услуги.
    canonical = canonical_nominal_name(value)
    return tuple(int(item) for item in re.findall(r"\d+", canonical))


def is_active_service(service: dict[str, Any]) -> bool:
    # Пропускает услугу только при явном признаке неактивности в исходном каталоге InterHub.
    raw = service.get("raw") if isinstance(service.get("raw"), dict) else {}
    value = raw.get("active")
    if value is None:
        return True
    return str(value).strip().lower() not in {"0", "false", "no", "off", "inactive"}


def collect_service_nominals(service: dict[str, Any]) -> list[dict[str, Any]]:
    # Берёт активные номиналы из того же поля каталога, которое использует форма платежа.
    nominal_field = next(
        (field for field in service.get("fields") or [] if str(field.get("name") or "").strip().casefold() == "nominal"),
        None,
    )
    if not nominal_field:
        return []
    nominals: list[dict[str, Any]] = []
    for item in nominal_field.get("value_list") or []:
        if not isinstance(item, dict) or item.get("active") is False:
            continue
        nominal_id = item.get("id")
        title = str(item.get("title") or item.get("name") or "").strip()
        if nominal_id is None or not title:
            continue
        nominals.append({"id": nominal_id, "title": title})
    return nominals


def normalize_detail_payload(payload: Any) -> list[dict[str, Any]]:
    # Поддерживает прямой список из нового метода и типовые обёртки data/items/result.
    raw_items: list[dict[str, Any]] = []

    def collect(value: Any) -> None:
        if isinstance(value, list):
            for item in value:
                collect(item)
            return
        if not isinstance(value, dict):
            return
        if value.get("name") is not None and value.get("count") is not None:
            raw_items.append(value)
            return
        for key in ("data", "items", "result"):
            nested = value.get(key)
            if isinstance(nested, (list, dict)):
                collect(nested)

    collect(payload)
    items: list[dict[str, Any]] = []
    for item in raw_items:
        try:
            numeric_count = float(item.get("count"))
            count: int | float | None = int(numeric_count) if numeric_count.is_integer() else numeric_count
        except (TypeError, ValueError):
            count = None
        items.append({"name": str(item.get("name") or "").strip(), "count": count})
    return items


def build_detail_loader(
    *,
    api_url: str,
    token: str,
    timeout_sec: int,
    ssl_verify: bool,
    ca_cert_path: str,
    proxy_url: str = "",
    detail_path: str = "/api/agent/service/detail",
) -> Callable[[int], dict[str, Any]]:
    # Создаёт отдельный GET-клиент нового метода, не меняя основной InterHub-сервис приложения.
    def get_service_detail(service_id: int) -> dict[str, Any]:
        try:
            cleaned_id = int(service_id)
        except (TypeError, ValueError) as exc:
            raise HTTPException(422, "InterHub service ID must be a positive integer") from exc
        if cleaned_id <= 0:
            raise HTTPException(422, "InterHub service ID must be a positive integer")
        if not str(api_url or "").strip():
            raise HTTPException(500, "InterHub API URL is not configured")
        if not str(token or "").strip():
            raise HTTPException(500, "InterHub token is not configured")

        separator = "&" if "?" in detail_path else "?"
        url = api_url.rstrip("/") + detail_path + separator + urllib.parse.urlencode({"id": cleaned_id})
        request = urllib.request.Request(
            url,
            headers={"Accept": "application/json", "Content-Type": "application/json", "token": str(token).strip()},
            method="GET",
        )
        context = ssl.create_default_context(cafile=ca_cert_path or None) if ssl_verify else ssl._create_unverified_context()
        try:
            if str(proxy_url or "").strip():
                proxy_handler = urllib.request.ProxyHandler({"http": proxy_url, "https": proxy_url})
                opener = urllib.request.build_opener(proxy_handler, urllib.request.HTTPSHandler(context=context))
                response_context = opener.open(request, timeout=max(5, int(timeout_sec or 20)))
            else:
                response_context = urllib.request.urlopen(request, timeout=max(5, int(timeout_sec or 20)), context=context)
            with response_context as response:
                raw = response.read() or b"[]"
            try:
                payload = json.loads(raw.decode("utf-8"))
            except (UnicodeDecodeError, json.JSONDecodeError) as exc:
                raise HTTPException(502, "InterHub service/detail returned invalid JSON") from exc
            return {"items": normalize_detail_payload(payload), "raw": payload}
        except urllib.error.HTTPError as exc:
            details = (exc.read() or b"").decode("utf-8", errors="ignore")
            suffix = f". {details}" if details else ""
            raise HTTPException(exc.code, f"InterHub {detail_path} failed: {exc.reason}{suffix}") from exc
        except urllib.error.URLError as exc:
            raise HTTPException(502, f"InterHub {detail_path} failed: {exc.reason}") from exc
        except (TimeoutError, socket.timeout) as exc:
            raise HTTPException(504, f"InterHub {detail_path} timed out") from exc

    return get_service_detail


def stock_row(
    service: dict[str, Any],
    nominal: dict[str, Any] | None,
    provider_name: str,
    count: int | float | None,
    match_status: str,
) -> dict[str, Any]:
    # Собирает единый формат строки, чтобы Excel не зависел от формы ответа поставщика.
    return {
        "service_id": service.get("service_id"),
        "service_title": service.get("title"),
        "category": service.get("category"),
        "service_type": service.get("type"),
        "nominal_id": nominal.get("id") if nominal else None,
        "nominal_title": nominal.get("title") if nominal else "",
        "provider_name": provider_name,
        "count": count,
        "match_status": match_status,
    }


def match_service_stock(service: dict[str, Any], detail_items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    # Сначала ищет точное имя, затем безопасные варианты только при единственной паре внутри услуги.
    local_items = collect_service_nominals(service)
    matched_local: dict[int, tuple[int, str]] = {}
    matched_provider: set[int] = set()

    def match_unique(key_builder: Callable[[Any], Any], status: str) -> None:
        # Назначает пару только когда ключ уникален с обеих сторон и не создаёт спорного соответствия.
        local_groups: dict[Any, list[int]] = {}
        provider_groups: dict[Any, list[int]] = {}
        for local_index, nominal in enumerate(local_items):
            if local_index in matched_local:
                continue
            key = key_builder(nominal.get("title"))
            if key:
                local_groups.setdefault(key, []).append(local_index)
        for provider_index, item in enumerate(detail_items):
            if provider_index in matched_provider:
                continue
            key = key_builder(item.get("name"))
            if key:
                provider_groups.setdefault(key, []).append(provider_index)
        for key, local_indexes in local_groups.items():
            provider_indexes = provider_groups.get(key, [])
            if len(local_indexes) == 1 and len(provider_indexes) == 1:
                local_index = local_indexes[0]
                provider_index = provider_indexes[0]
                matched_local[local_index] = (provider_index, status)
                matched_provider.add(provider_index)

    match_unique(normalize_nominal_name, "Совпало по имени")
    match_unique(canonical_nominal_name, "Совпало после нормализации")
    match_unique(nominal_number_signature, "Совпало по числам в имени — проверить")

    rows: list[dict[str, Any]] = []
    for local_index, nominal in enumerate(local_items):
        matched = matched_local.get(local_index)
        if matched:
            provider_index, status = matched
            provider = detail_items[provider_index]
            rows.append(stock_row(service, nominal, str(provider.get("name") or ""), provider.get("count"), status))
            continue

        # Если уникальная пара не нашлась, показываем возможные дубли, но не объявляем их точным совпадением.
        local_key = canonical_nominal_name(nominal.get("title"))
        candidates = [
            (provider_index, item)
            for provider_index, item in enumerate(detail_items)
            if provider_index not in matched_provider and canonical_nominal_name(item.get("name")) == local_key
        ]
        if candidates:
            for provider_index, _ in candidates:
                matched_provider.add(provider_index)
            counts = [item.get("count") for _, item in candidates if isinstance(item.get("count"), (int, float))]
            provider_names = " | ".join(dict.fromkeys(str(item.get("name") or "") for _, item in candidates))
            rows.append(stock_row(service, nominal, provider_names, sum(counts) if counts else None, "Неоднозначное имя — проверить"))
        else:
            rows.append(stock_row(service, nominal, "", None, "Нет строки в service/detail"))

    for provider_index, item in enumerate(detail_items):
        if provider_index in matched_provider:
            continue
        rows.append(stock_row(service, None, str(item.get("name") or ""), item.get("count"), "Нет номинала в нашем каталоге"))

    if not rows:
        rows.append(stock_row(service, None, "", None, "У услуги нет номиналов и строк остатков"))
    return rows


def collect_stock(
    client,
    services: list[dict[str, Any]],
    *,
    delay_ms: int,
    progress=None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    # Вызывает service/detail один раз на услугу и продолжает формировать отчёт после единичной ошибки.
    active_services = [service for service in services if is_active_service(service)]
    stock_rows: list[dict[str, Any]] = []
    api_rows: list[dict[str, Any]] = []
    for index, service in enumerate(active_services, start=1):
        if index > 1 and delay_ms > 0:
            time.sleep(delay_ms / 1000)
        try:
            detail = client.get_service_detail(int(service["service_id"]))
            items = detail.get("items") if isinstance(detail.get("items"), list) else []
            service_rows = match_service_stock(service, items)
            api_rows.append({"service_id": service.get("service_id"), "service_title": service.get("title"), "success": True, "message": "", "response": detail.get("raw")})
        except Exception as exc:  # noqa: BLE001 - ошибка одной услуги должна остаться в итоговом файле
            message = str(getattr(exc, "detail", exc) or exc)
            service_rows = [stock_row(service, None, "", None, f"Ошибка service/detail: {message}")]
            api_rows.append({"service_id": service.get("service_id"), "service_title": service.get("title"), "success": False, "message": message, "response": {}})
        stock_rows.extend(service_rows)
        if progress:
            progress(index, len(active_services), service, service_rows)
    return stock_rows, api_rows


def json_text(value: Any) -> str:
    # Сохраняет полный ответ поставщика для разбора спорных совпадений без повторного запроса.
    try:
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    except (TypeError, ValueError):
        return str(value or "")


def build_report(stock_rows: list[dict[str, Any]], api_rows: list[dict[str, Any]], *, generated_at: datetime) -> Workbook:
    # Создаёт сводку, рабочий лист остатков и технический лист исходных ответов.
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    # Убираем timezone без перевода, потому что время уже передано в московской зоне для отчёта.
    excel_generated_at = generated_at.replace(tzinfo=None) if generated_at.tzinfo else generated_at
    summary_rows = [
        ("Остатки InterHub", "service/list → service/detail, без calculate/check/pay"),
        ("Сформировано", excel_generated_at),
        ("Активных услуг опрошено", len(api_rows)),
        ("Номиналов сопоставлено", sum(str(row["match_status"]).startswith("Совпало") for row in stock_rows)),
        ("Из них по числам — проверить", sum(row["match_status"] == "Совпало по числам в имени — проверить" for row in stock_rows)),
        ("Нет строки в service/detail", sum(row["match_status"] == "Нет строки в service/detail" for row in stock_rows)),
        ("Только у поставщика", sum(row["match_status"] == "Нет номинала в нашем каталоге" for row in stock_rows)),
        ("Неоднозначных совпадений", sum(row["match_status"] == "Неоднозначное имя — проверить" for row in stock_rows)),
        ("Ошибок service/detail", sum(not row["success"] for row in api_rows)),
    ]
    for item in summary_rows:
        summary.append(item)
    style_summary_sheet(summary)

    stock = workbook.create_sheet("Остатки")
    stock.append(STOCK_HEADERS)
    for row in stock_rows:
        stock.append([
            row.get("service_id"), row.get("service_title"), row.get("category"), row.get("service_type"),
            row.get("nominal_id"), row.get("nominal_title"), row.get("provider_name"), row.get("count"),
            row.get("match_status"),
        ])
    style_stock_sheet(stock)

    responses = workbook.create_sheet("Ответы service detail")
    responses.append(["ID услуги", "Услуга", "Успешно", "Ошибка", "Полный ответ service/detail (JSON)"])
    for row in api_rows:
        responses.append([
            row.get("service_id"), row.get("service_title"), row.get("success"), row.get("message"),
            json_text(row.get("response")),
        ])
    style_response_sheet(responses)
    return workbook


def apply_header_style(sheet) -> None:
    # Выделяет заголовок и закрепляет его для удобной фильтрации длинного отчёта.
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def style_summary_sheet(sheet) -> None:
    # Оформляет компактную сводку и сохраняет дату настоящим значением Excel.
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    sheet["B2"].number_format = "yyyy-mm-dd hh:mm:ss"
    sheet.column_dimensions["A"].width = 31
    sheet.column_dimensions["B"].width = 54


def style_stock_sheet(sheet) -> None:
    # Подсвечивает точные совпадения, расхождения и ошибки разными цветами.
    apply_header_style(sheet)
    fills = {
        "Совпало по имени": PatternFill("solid", fgColor="C6EFCE"),
        "Совпало после нормализации": PatternFill("solid", fgColor="C6EFCE"),
        "Совпало по числам в имени — проверить": PatternFill("solid", fgColor="FCE4D6"),
        "Нет строки в service/detail": PatternFill("solid", fgColor="FFEB9C"),
        "Нет номинала в нашем каталоге": PatternFill("solid", fgColor="DDEBF7"),
        "Неоднозначное имя — проверить": PatternFill("solid", fgColor="FCE4D6"),
    }
    for row_index in range(2, sheet.max_row + 1):
        sheet.cell(row_index, 8).number_format = "#,##0.##"
        status_cell = sheet.cell(row_index, 9)
        status_cell.fill = fills.get(str(status_cell.value or ""), PatternFill("solid", fgColor="FFC7CE"))
        status_cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in range(1, sheet.max_column + 1):
        longest = max((len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1)), default=10)
        sheet.column_dimensions[get_column_letter(column)].width = min(max(longest + 2, 12), 44)


def style_response_sheet(sheet) -> None:
    # Делает технический JSON читаемым и не раздувает ширину основного листа.
    apply_header_style(sheet)
    for row_index in range(2, sheet.max_row + 1):
        for column in (4, 5):
            sheet.cell(row_index, column).alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in {"A": 14, "B": 32, "C": 12, "D": 42, "E": 72}.items():
        sheet.column_dimensions[column].width = width


def build_client_from_env():
    # Повторно использует штатную загрузку каталога, добавляя только изолированный вызов service/detail.
    api_url = os.getenv("INTERHUB_API_URL", "")
    token = os.getenv("INTERHUB_TOKEN", "")
    timeout_sec = int(os.getenv("INTERHUB_TIMEOUT_SEC", "20") or "20")
    ssl_verify = parse_bool(os.getenv("INTERHUB_SSL_VERIFY"), True)
    ca_cert_path = os.getenv("INTERHUB_CA_CERT_PATH", "")
    proxy_url = os.getenv("INTERHUB_PROXY_URL", "")
    catalog_client = build_interhub_service(
        HTTPException=HTTPException,
        interhub_api_url=api_url,
        interhub_token=token,
        timeout_sec=timeout_sec,
        ssl_verify=ssl_verify,
        ca_cert_path=ca_cert_path,
        proxy_url=proxy_url,
        calculate_path=os.getenv("INTERHUB_CALCULATE_PATH", "/api/agent/payment/check/calculate"),
        check_path=os.getenv("INTERHUB_CHECK_PATH", "/api/agent/payment/check"),
        pay_path=os.getenv("INTERHUB_PAY_PATH", "/api/agent/payment/pay"),
        check_status_path=os.getenv("INTERHUB_CHECK_STATUS_PATH", "/api/agent/payment/check_status"),
        deposit_path=os.getenv("INTERHUB_DEPOSIT_PATH", "/api/agent/deposit"),
    )
    detail_loader = build_detail_loader(
        api_url=api_url,
        token=token,
        timeout_sec=timeout_sec,
        ssl_verify=ssl_verify,
        ca_cert_path=ca_cert_path,
        proxy_url=proxy_url,
        detail_path=os.getenv("INTERHUB_SERVICE_DETAIL_PATH", "/api/agent/service/detail"),
    )
    return SimpleNamespace(get_services=catalog_client.get_services, get_service_detail=detail_loader)


def parse_args() -> argparse.Namespace:
    # Оставляет фильтры необязательными: обычный запуск обходит весь активный каталог.
    parser = argparse.ArgumentParser(description="Выгрузить остатки InterHub через service/detail")
    parser.add_argument("--output", type=Path, help="Путь к итоговому .xlsx")
    parser.add_argument("--delay-ms", type=int, default=250, help="Пауза между запросами service/detail")
    parser.add_argument("--limit-services", type=int, default=0, help="Проверить только первые N активных услуг")
    parser.add_argument("--service-id", action="append", type=int, default=[], help="Проверять только указанный ID услуги")
    return parser.parse_args()


def main() -> int:
    # Загружает каталог, собирает остатки и оставляет готовый Excel в указанном каталоге сервера.
    load_dotenv(ROOT_DIR / ".env.dev", override=False)
    args = parse_args()
    if args.delay_ms < 0 or args.limit_services < 0:
        print("--delay-ms и --limit-services не могут быть отрицательными", file=sys.stderr)
        return 2

    client = build_client_from_env()
    services = [service for service in client.get_services() if is_active_service(service)]
    if args.service_id:
        allowed_ids = set(args.service_id)
        services = [service for service in services if service.get("service_id") in allowed_ids]
    if args.limit_services:
        services = services[:args.limit_services]
    if not services:
        print("В каталоге нет подходящих активных услуг", file=sys.stderr)
        return 1

    started_at = datetime.now(ZoneInfo("Europe/Moscow"))

    def print_progress(index: int, total: int, service: dict[str, Any], service_rows: list[dict[str, Any]]) -> None:
        # Показывает прогресс и число строк, не раскрывая полные ответы поставщика.
        print(f"[{index}/{total}] {service.get('service_id')} · {service.get('title')}: {len(service_rows)} строк")

    stock_rows, api_rows = collect_stock(client, services, delay_ms=args.delay_ms, progress=print_progress)
    workbook = build_report(stock_rows, api_rows, generated_at=started_at)
    output = args.output or Path.cwd() / f"interhub-service-stock-{started_at:%Y%m%d-%H%M%S}.xlsx"
    output = output.expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"Готово: {output}")
    print(f"Услуг опрошено: {len(api_rows)}; строк остатков: {len(stock_rows)}; calculate/check/pay не вызывались")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
