"""История и Excel читают журнал без поставщика и без секретных результатов."""
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
import unittest
from unittest.mock import MagicMock

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from domains.airpay_api import mount_airpay_routes
from domains.airpay_history import HistoryFilters, history_where, export_history, EXPORT_LIMIT, XLSX_TYPE
from domains.airpay_repository import AirpayRepository


def saved_row(**changes):
    # Длинный ID и похожее на формулу поле проверяют точность и безопасность Excel.
    return dict(agent_transaction_id=1645129032053759724, service_title='=HYPERLINK("bad")',
                service_id='A0217', account='seller@example.test', state='paid', purchase_kind='voucher',
                amount=Decimal('654.03'), currency='RUB', provider_transaction_id='999999999999999999',
                service_snapshot={}, created_at=datetime(2026, 9, 22, 22, 0, tzinfo=timezone.utc),
                pin_code='DO-NOT-EXPORT', pin_ciphertext='ENCRYPTED', request_payload={'password': 'SECRET'},
                voucher_response={'pinCode': 'DO-NOT-EXPORT'}, **changes)


class AirpayHistoryTests(unittest.TestCase):
    def setUp(self):
        # Любой случайный вызов метода провайдера падает; только репозиторий подменён.
        self.service = MagicMock()
        self.service.payments_enabled = False
        for name in ('pay', 'check', 'voucher', 'get_balance', 'get_service', 'get_services'):
            getattr(self.service, name).side_effect = AssertionError('Provider forbidden')
        self.repo = MagicMock()
        self.repo.history.return_value = [saved_row()]
        app = FastAPI()
        mount_airpay_routes(app, get_current_user=lambda: {'role': 'owner', 'username': 'alice'}, service=self.service,
                            repository=self.repo, offline=True)
        self.client = TestClient(app)

    def test_filters_are_shared_with_export_and_page_limits_are_ignored_for_export(self):
        # Поиск Excel не обрезается текущей страницей UI и не раскрывает ваучер.
        query = '?q=A0217&kind=voucher&status=awaiting_voucher&date_from=2026-09-01&date_to=2026-09-23&limit=1&offset=20'
        response = self.client.get('/integrations/airpay/transactions' + query)
        self.assertEqual(response.status_code, 200)
        filters = self.repo.history.call_args.args[3]
        self.assertNotIn('DO-NOT-EXPORT', response.text)
        response = self.client.get('/integrations/airpay/transactions/export' + query)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['content-type'], XLSX_TYPE)
        self.repo.history.assert_called_with('alice', EXPORT_LIMIT + 1, 0, filters)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet['A2'].value, '1645129032053759724')
        self.assertEqual(sheet['B2'].value, '2026-09-23 01:00:00')
        self.assertEqual(sheet['C2'].data_type, 's')
        self.assertEqual(sheet['H2'].value, '654.03')
        self.assertEqual(sheet['J2'].value, '999999999999999999')
        contents = str(list(sheet.values))
        for secret in ('DO-NOT-EXPORT', 'ENCRYPTED', 'SECRET', 'pinCode'):
            self.assertNotIn(secret, contents)
        self.assertEqual(response.headers['cache-control'], 'no-store')

    def test_bad_filters_fail_before_repository(self):
        # Неизвестные состояния и некорректные даты не превращаются в поиск без ограничений.
        for query in ('status=unknown', 'kind=oops', 'q=' + 'a' * 201, 'date_from=bad',
                      'date_from=2026-09-24&date_to=2026-09-23', 'date_to=9999-12-31'):
            for path in ('transactions', 'transactions/export'):
                self.assertEqual(self.client.get('/integrations/airpay/' + path + '?' + query).status_code, 422)
        self.repo.history.assert_not_called()

    def test_sql_scopes_owner_and_escapes_search_with_moscow_day_bounds(self):
        # Шаблоны SQL не позволяют поиску захватить чужие строки или расшириться через % и _.
        filters = HistoryFilters(q="%' OR 1=1 --_\\", date_from='2026-09-23', date_to='2026-09-23', kind='voucher', status='processing')
        where, params = history_where('alice', filters)
        self.assertTrue(where.startswith('created_by=%s AND ('))
        self.assertNotIn('OR 1=1', where)
        self.assertEqual(params[0], 'alice')
        self.assertEqual(params[1], "%\\%' OR 1=1 --\\_\\\\%")
        self.assertEqual(params[-3].astimezone(timezone.utc).isoformat(), '2026-09-22T21:00:00+00:00')
        self.assertEqual(params[-2].astimezone(timezone.utc).isoformat(), '2026-09-23T21:00:00+00:00')
        connection = MagicMock()
        connection.return_value.__enter__.return_value.execute.return_value.fetchall.return_value = []
        AirpayRepository(connection).history('alice', 21, 40, filters)
        query, args = connection.return_value.__enter__.return_value.execute.call_args.args
        self.assertIn(where, query)
        self.assertEqual(args, (*params, 21, 40))

    def test_paid_missing_code_is_not_completed_and_attention_is_separate(self):
        # Обе формы хранения старого кода учитываются в состоянии результата.
        waiting, _ = history_where('alice', HistoryFilters(status='awaiting_voucher'))
        complete, _ = history_where('alice', HistoryFilters(status='completed'))
        attention, _ = history_where('alice', HistoryFilters(status='requires_attention'))
        self.assertIn("purchase_kind='voucher' AND NOT", waiting)
        self.assertIn('pin_ciphertext IS NOT NULL', waiting)
        self.assertIn("COALESCE(pin_code, '') <> ''", waiting)
        self.assertIn("state='paid' AND (purchase_kind='topup' OR", complete)
        self.assertIn('requires_attention=true', attention)

    def test_export_refuses_truncation_and_allows_empty_result(self):
        # Успешная выгрузка всегда полная; пустой журнал содержит заголовки.
        self.repo.history.return_value = [None] * (EXPORT_LIMIT + 1)
        with self.assertRaises(HTTPException) as caught:
            export_history(self.repo, 'alice', HistoryFilters())
        self.assertEqual(caught.exception.status_code, 422)
        self.repo.history.return_value = []
        sheet = load_workbook(BytesIO(export_history(self.repo, 'alice', HistoryFilters()).body)).active
        self.assertEqual(sheet.max_row, 1)

    def test_export_requires_owner(self):
        # Отдельный download URL не расширяет доступ продавца к журналу владельца.
        app = FastAPI()
        mount_airpay_routes(app, get_current_user=lambda: {'role': 'seller', 'username': 'alice'}, service=self.service, repository=self.repo)
        self.assertEqual(TestClient(app).get('/integrations/airpay/transactions/export').status_code, 403)
        self.repo.history.assert_not_called()
