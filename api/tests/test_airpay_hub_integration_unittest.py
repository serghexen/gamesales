"""CRM → настоящий Airpay ASGI Hub → подменённый поставщик; никаких сетевых оплат."""
import io
import json
import os
from pathlib import Path
import sys
import unittest
import urllib.error
from urllib.parse import urlsplit
from unittest.mock import MagicMock, patch
from uuid import uuid4

from fastapi import FastAPI
from fastapi.testclient import TestClient
from domains.airpay_hub_client import mount_airpay_hub_proxy
from tests.test_airpay_purchase_unittest import MemoryJournal


HUB_SOURCE = os.getenv('AIRPAY_HUB_TEST_SOURCE', '')


@unittest.skipUnless(HUB_SOURCE, 'Укажите AIRPAY_HUB_TEST_SOURCE для межрепозиторного теста без сети')
class AirpayHubIntegrationTests(unittest.TestCase):
    def setUp(self):
        # Загружаем приложение Hub из явного пути, без чтения .env и запуска серверов/worker.
        hub_api = str(Path(HUB_SOURCE) / 'api')
        sys.path.insert(0, hub_api)
        self.addCleanup(lambda: sys.path.remove(hub_api))
        from hub.airpay_app import create_airpay_app
        self.journal = MemoryJournal()
        self.provider = MagicMock(payments_enabled=True)
        self.provider.get_service.return_value = {
            'service_id': 'A0217', 'title': 'Test voucher', 'inputs': [], 'displays': [], 'fixed_payment': True}
        self.provider.get_balance.return_value = {'configured': True, 'balance': 10000, 'currency': 'RUB'}
        self.provider.check.side_effect = lambda payload: {
            'result': 0, 'agentTransactionId': payload['agentTransactionId'], 'transactionId': 70, 'fixedPrice': 12.5}
        self.provider.pay.side_effect = lambda payload: {
            'result': 0, 'agentTransactionId': payload['agentTransactionId'], 'transactionId': 71}
        self.provider.get_voucher.side_effect = lambda payload: {
            'result': 0, 'agentTransactionId': payload['agentTransactionId'], 'transactionId': 71,
            'displays': {'pinCode': 'FAKE-RESULT-NOT-A-REAL-VOUCHER'}}
        env = {'DATABASE_URL': 'unused', 'SUPPLIER_HUB_DATA_SECRET': 'd' * 32,
               'SUPPLIER_HUB_CLIENTS_JSON': json.dumps({'crm': 'k' * 32}),
               'SUPPLIER_HUB_PURCHASES_ENABLED': 'true', 'AIRPAY_PAYMENTS_ENABLED': 'true'}
        with patch('hub.airpay_app.build_airpay_service', return_value=self.provider), \
             patch('hub.airpay_app.AirpayRepository', return_value=self.journal), \
             patch('hub.airpay_app.AirpayJobStore', return_value=None):
            self.hub = TestClient(create_airpay_app(env, MagicMock()))
        self.addCleanup(self.hub.close)
        self.calls, self.lose_pay_response = [], False
        opener = MagicMock()
        opener.open.side_effect = self.deliver_in_memory
        patcher = patch('domains.airpay_hub_client.urllib.request.build_opener', return_value=opener)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.owner = 'owner'
        app = FastAPI()
        mount_airpay_hub_proxy(app, get_current_user=lambda: {'role': 'owner', 'username': self.owner},
            environ={'AIRPAY_HUB_URL': 'http://airpay-api:8011', 'AIRPAY_HUB_CLIENT_ID': 'crm', 'AIRPAY_HUB_CLIENT_KEY': 'k' * 32})
        self.crm = TestClient(app)
        self.addCleanup(self.crm.close)

    def deliver_in_memory(self, request, timeout):
        # urllib подменён адаптером TestClient: исходящий HTTP не создаётся даже для метода pay.
        parsed = urlsplit(request.full_url)
        self.assertEqual(parsed.netloc, 'airpay-api:8011')
        self.assertNotIn('Authorization', request.headers)
        self.calls.append((request.method, parsed.path))
        response = self.hub.request(request.method, parsed.path + ('?' + parsed.query if parsed.query else ''),
                                    content=request.data, headers=dict(request.header_items()))
        if self.lose_pay_response and parsed.path.endswith('/pay'):
            self.lose_pay_response = False
            raise urllib.error.URLError('Simulated lost response after saved payment')
        if response.status_code >= 400:
            raise urllib.error.HTTPError(request.full_url, response.status_code, '', response.headers, io.BytesIO(response.content))
        wrapped = MagicMock()
        wrapped.status, wrapped.headers = response.status_code, response.headers
        wrapped.read.side_effect = io.BytesIO(response.content).read
        wrapped.__enter__.return_value = wrapped
        return wrapped

    def prepare_and_check(self):
        # Один стабильный ключ проходит оба приложения, исходные ID назначает только Hub.
        payload = {'service_id': 'A0217', 'fields': {'account': 'test@example.invalid'}, 'preparation_key': str(uuid4())}
        response = self.crm.post('/integrations/airpay/prepare', json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        draft = response.json()
        again = self.crm.post('/integrations/airpay/prepare', json=payload)
        self.assertEqual(again.json()['agent_transaction_id'], draft['agent_transaction_id'])
        changed = self.crm.post('/integrations/airpay/prepare', json={**payload, 'fields': {'account': 'changed@example.invalid'}})
        self.assertEqual(changed.status_code, 409)
        checked = self.crm.post('/integrations/airpay/check', json={'preparation_token': draft['preparation_token']})
        self.assertEqual(checked.status_code, 200, checked.text)
        self.assertTrue(checked.json()['purchase_ready'])
        self.provider.pay.assert_not_called()
        self.provider.get_voucher.assert_not_called()
        return draft['agent_transaction_id']

    def test_full_route_saves_result_and_reveals_only_by_separate_action(self):
        # Оплата и получение кода проверяются на MagicMock, а оба API и сериализация настоящие.
        transaction_id = self.prepare_and_check()
        path = '/integrations/airpay/transactions/' + transaction_id
        paid = self.crm.post(path + '/pay', json={'confirmed_amount': '12.50'})
        self.assertEqual(paid.status_code, 200, paid.text)
        self.assertEqual(paid.json()['payment_state'], 'paid')
        self.assertEqual(paid.json()['result_state'], 'pending')
        self.assertFalse(paid.json()['completed'])
        self.provider.get_voucher.assert_not_called()
        self.assertEqual(self.crm.get(path).json()['result_state'], 'pending')
        result = self.crm.post(path + '/voucher')
        self.assertEqual(result.status_code, 200, result.text)
        self.assertTrue(result.json()['completed'])
        self.assertNotIn('FAKE-RESULT', result.text)
        self.assertNotIn('FAKE-RESULT', self.crm.get('/integrations/airpay/transactions').text)
        self.crm.post(path + '/voucher')
        self.crm.post(path + '/pay', json={'confirmed_amount': '12.50'})
        self.provider.pay.assert_called_once()
        self.provider.get_voucher.assert_called_once()
        revealed = self.crm.post(path + '/result')
        self.assertEqual(revealed.json()['value'], 'FAKE-RESULT-NOT-A-REAL-VOUCHER')
        self.assertEqual(self.journal.accesses, [('["crm","owner"]', int(transaction_id))])
        self.owner = 'other-owner'
        self.assertEqual(self.crm.get(path).status_code, 404)
        self.assertEqual(self.crm.post(path + '/result').status_code, 404)

    def test_lost_http_response_reads_same_hub_operation_without_new_pay(self):
        # Потеря ответа CRM не означает отказ оплаты: чтение возвращает сохранённый успех Hub.
        transaction_id = self.prepare_and_check()
        path = '/integrations/airpay/transactions/' + transaction_id
        self.lose_pay_response = True
        response = self.crm.post(path + '/pay', json={'confirmed_amount': '12.50'})
        self.assertEqual(response.status_code, 503)
        saved = self.crm.get(path)
        self.assertEqual(saved.json()['payment_state'], 'paid')
        self.assertEqual(saved.json()['result_state'], 'pending')
        self.provider.pay.assert_called_once()
        self.provider.get_voucher.assert_not_called()
        self.assertEqual(len(self.journal.rows), 1)

    def test_export_crosses_both_real_apis_without_any_supplier_call(self):
        # Сохранённая операция проходит авторизацию обоих API и возвращается настоящим XLSX.
        from openpyxl import load_workbook
        from tests.test_airpay_history_unittest import saved_row
        row = saved_row(created_by='["crm","owner"]')
        self.journal.rows[row['agent_transaction_id']] = row
        response = self.crm.get('/integrations/airpay/transactions/export')
        self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else '')
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet['A2'].value, '1645129032053759724')
        self.assertEqual(sheet['K2'].value, 'Код сохранён')
        self.assertEqual(self.calls, [('GET', '/integrations/airpay/transactions/export')])
        self.assertEqual(self.provider.mock_calls, [])
        self.owner = 'other-owner'
        response = self.crm.get('/integrations/airpay/transactions/export')
        self.assertEqual(load_workbook(io.BytesIO(response.content)).active.max_row, 1)
