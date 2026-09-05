import json
import sys
import unittest
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook


SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from export_interhub_service_stock import (  # noqa: E402
    build_detail_loader,
    build_report,
    collect_stock,
    match_service_stock,
    normalize_detail_payload,
)


class _Response:
    def __init__(self, payload):
        # Имитирует JSON-ответ urllib без обращения к поставщику.
        self.raw = json.dumps(payload).encode("utf-8")

    def read(self):
        # Возвращает подготовленные байты ответа.
        return self.raw

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_value, traceback):
        return False


class _FakeClient:
    def __init__(self):
        # Хранит ID вызванных услуг, чтобы доказать один detail-запрос на каждую активную услугу.
        self.detail_calls = []

    def get_service_detail(self, service_id):
        # Возвращает остатки первой услуги и ошибку второй, не останавливая общий отчёт.
        self.detail_calls.append(service_id)
        if service_id == 11:
            raw = [{"name": " 60\u00a0CP ", "count": 10}, {"name": "1580 CP", "count": 5}, {"name": "Extra", "count": 2}]
            return {"items": normalize_detail_payload(raw), "raw": raw}
        raise RuntimeError("detail unavailable")


class InterHubServiceStockExportScriptTests(unittest.TestCase):
    def setUp(self):
        # Готовит активную, ошибочную и явно выключенную услуги с номиналами каталога.
        self.services = [
            {
                "service_id": 11,
                "title": "po_Arena Breakout - Global",
                "category": "Games",
                "type": "VOUCHER",
                "fields": [{"name": "nominal", "value_list": [{"id": 101, "title": "60 CP"}, {"id": 102, "title": "1580 cp"}, {"id": 103, "title": "3200 CP"}]}],
                "raw": {"active": True},
            },
            {
                "service_id": 12,
                "title": "Broken service",
                "category": "Games",
                "type": "VOUCHER",
                "fields": [{"name": "nominal", "value_list": [{"id": 201, "title": "100 UC"}]}],
                "raw": {"active": True},
            },
            {"service_id": 13, "title": "Disabled", "fields": [], "raw": {"active": False}},
        ]

    def test_detail_loader_sends_one_get_with_service_id_and_token(self):
        # Проверяет контракт нового метода и отсутствие тела запроса.
        loader = build_detail_loader(
            api_url="https://api.interhub.uz",
            token="test-token",
            timeout_sec=20,
            ssl_verify=False,
            ca_cert_path="",
        )
        payload = [{"name": "60 CP", "count": 10}]
        with patch("export_interhub_service_stock.urllib.request.urlopen", return_value=_Response(payload)) as urlopen_mock:
            result = loader(9931)

        request = urlopen_mock.call_args.args[0]
        self.assertEqual(request.method, "GET")
        self.assertIsNone(request.data)
        self.assertEqual(request.get_header("Token"), "test-token")
        self.assertEqual(request.full_url, "https://api.interhub.uz/api/agent/service/detail?id=9931")
        self.assertEqual(result["items"], [{"name": "60 CP", "count": 10}])

    def test_matching_uses_normalized_name_and_keeps_both_kinds_of_mismatch(self):
        # Сопоставляет регистр и пробелы, но не скрывает номиналы, отсутствующие с любой стороны.
        detail = normalize_detail_payload([{"name": " 60\u00a0cp ", "count": 10}, {"name": "Extra", "count": 2}])
        rows = match_service_stock(self.services[0], detail)

        by_nominal = {row["nominal_id"]: row for row in rows if row["nominal_id"] is not None}
        self.assertEqual(by_nominal[101]["count"], 10)
        self.assertEqual(by_nominal[101]["match_status"], "Совпало по имени")
        self.assertEqual(by_nominal[102]["match_status"], "Нет строки в service/detail")
        self.assertTrue(any(row["provider_name"] == "Extra" and row["match_status"] == "Нет номинала в нашем каталоге" for row in rows))

    def test_matching_uses_unique_number_signature_only_inside_one_service(self):
        # Сопоставляет разные оформления одного номинала и помечает резервный результат для проверки.
        detail = normalize_detail_payload([{"name": "gf_PUBG MOBILE 12,000 + 4,200 UC (GLOBAL)", "count": 2}])
        service = {
            **self.services[0],
            "fields": [{"name": "nominal", "value_list": [{"id": 104, "title": "12.000 + 4.200 UC"}]}],
        }
        rows = match_service_stock(service, detail)

        self.assertEqual(rows[0]["count"], 2)
        self.assertEqual(rows[0]["match_status"], "Совпало по числам в имени — проверить")

    def test_collect_calls_detail_once_per_active_service_and_continues_after_error(self):
        # Доказывает лёгкий обход: один detail на услугу, без вызовов calculate/check/pay.
        client = _FakeClient()
        stock_rows, api_rows = collect_stock(client, self.services, delay_ms=0)

        self.assertEqual(client.detail_calls, [11, 12])
        self.assertEqual(len(api_rows), 2)
        self.assertTrue(api_rows[0]["success"])
        self.assertFalse(api_rows[1]["success"])
        self.assertTrue(any(row["match_status"].startswith("Ошибка service/detail") for row in stock_rows))

    def test_report_contains_stock_and_raw_response_sheets(self):
        # Проверяет листы, числовой остаток, фильтры и локальную дату Excel без timezone.
        client = _FakeClient()
        stock_rows, api_rows = collect_stock(client, self.services[:1], delay_ms=0)
        generated_at = datetime(2026, 8, 28, 20, 30, tzinfo=timezone(timedelta(hours=3)))
        workbook = build_report(stock_rows, api_rows, generated_at=generated_at)
        buffer = BytesIO()
        workbook.save(buffer)
        loaded = load_workbook(BytesIO(buffer.getvalue()))

        self.assertEqual(loaded.sheetnames, ["Сводка", "Остатки", "Ответы service detail"])
        self.assertEqual(loaded["Сводка"]["B3"].value, 1)
        self.assertIsNone(loaded["Сводка"]["B2"].value.tzinfo)
        self.assertEqual(loaded["Сводка"]["B2"].value, generated_at.replace(tzinfo=None))
        stock = loaded["Остатки"]
        self.assertEqual(stock["H2"].value, 10)
        self.assertEqual(stock["I2"].value, "Совпало по имени")
        self.assertEqual(stock.freeze_panes, "A2")
        self.assertEqual(stock.auto_filter.ref, stock.dimensions)
        responses = loaded["Ответы service detail"]
        self.assertIn('"count": 10', responses["E2"].value)


if __name__ == "__main__":
    unittest.main()
