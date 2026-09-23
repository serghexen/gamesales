import json
import unittest
from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, Mock, patch

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import BaseModel

from domains.interhub_api import mount_interhub_routes
from domains.interhub_price_cache import build_interhub_prices_xlsx, collect_price_targets
from domains.interhub_service import build_interhub_service
from domains.interhub_stock_cache import fetch_nominal_stock, match_stock_targets, parse_stock_count, save_stock_rows


def catalog():
    # Два сервиса проверяют, что detail вызывается по услуге, а calculate — по каждому номиналу.
    return [
        {'service_id': 7, 'title': 'Gift', 'type': 'VOUCHER', 'fields': [
            {'name': 'nominal', 'value_list': [{'id': 11, 'title': 'USD 10'}, {'id': 12, 'title': 'USD 20'}]}]},
        {'service_id': 8, 'title': 'Top up', 'type': 'TOP_UP_FIXED', 'fields': [
            {'name': 'nominal', 'value_list': [{'id': 13, 'title': '60 CP'}]}]},
    ]


class ApiModel(BaseModel):
    model_config = {'extra': 'allow'}


class StockCacheTests(unittest.TestCase):
    def test_live_stock_matches_only_inside_service_and_keeps_ambiguous_names_unknown(self):
        # Одинаковые имена разных услуг не смешиваем, а дубли в одной услуге не угадываем.
        services = catalog()
        detail = Mock(return_value=[{'name': 'USD 10.00', 'count': 4}])
        stock = fetch_nominal_stock(7, '11', get_services=lambda: services, get_detail=detail)
        self.assertEqual(stock['stock_count'], 4)
        self.assertEqual(stock['match_status'], 'normalized')
        detail.assert_called_once_with(7)
        services[0]['fields'][0]['value_list'].append({'id': 99, 'title': 'USD 10.00'})
        stock = fetch_nominal_stock(7, 11, get_services=lambda: services, get_detail=detail)
        self.assertIsNone(stock['stock_count'])
        self.assertEqual(stock['match_status'], 'ambiguous')

    def test_live_stock_catalog_failure_keeps_raw_reply_and_never_guesses_nominal(self):
        # Без актуального имени не используем прежний кэш, даже если detail вернул число.
        payload = [{'name': 'USD 10', 'count': 7}]
        for services in ([], catalog()):
            stock = fetch_nominal_stock(7, 999, get_services=lambda: services, get_detail=lambda _: payload)
            self.assertIsNone(stock['stock_count'])
            self.assertEqual(stock['match_status'], 'error')
            self.assertEqual(stock['provider_response'], payload)
            self.assertIsNotNone(stock['checked_at'].tzinfo)

    def test_zero_missing_errors_and_bad_counts_have_different_meanings(self):
        # Ошибка внутри HTTP 200 и неизвестный номинал не маскируются под нулевой запас.
        targets = collect_price_targets(catalog())[:2]
        rows = match_stock_targets(targets, [{'name': 'usd\u00a010', 'count': 0}])
        self.assertEqual(rows[0]['stock_count'], 0)
        self.assertEqual(rows[0]['match_status'], 'matched')
        self.assertIsNone(rows[1]['stock_count'])
        self.assertEqual(rows[1]['match_status'], 'missing')
        for payload in ({'success': False, 'status': -1, 'message': 'Error'}, {}, None):
            rows = match_stock_targets(targets, payload)
            self.assertTrue(all(row['match_status'] == 'error' and row['stock_count'] is None for row in rows))
            self.assertEqual(rows[0]['provider_response'], payload)
        self.assertEqual(match_stock_targets(targets, [])[0]['match_status'], 'missing')
        self.assertEqual(match_stock_targets(targets, [{'name': 'USD 10', 'count': 'oops'}])[0]['match_status'], 'invalid_count')
        for value in (None, True, False, -1, 1.2, 'NaN', 'Infinity', '9223372036854775808'):
            self.assertIsNone(parse_stock_count(value))
        self.assertEqual(parse_stock_count('12.00'), 12)

    def test_normalization_keeps_units_and_rejects_ambiguous_names(self):
        # Формат числа может различаться, но одинаковые цифры другой валюты не дают ложного остатка.
        target = collect_price_targets(catalog())[0]
        for local_name, provider_name in [('USD 1000', 'USD 1,000'), ('USD 10', 'USD 10.00'), ('1500+80 NC', '1500 + 80 NC')]:
            row = match_stock_targets([{**target, 'nominal_title': local_name}], [{'name': provider_name, 'count': 2}])[0]
            self.assertEqual(row['stock_count'], 2)
            self.assertEqual(row['match_status'], 'normalized')
        row = match_stock_targets([target], [{'name': 'EUR 10', 'count': 55}])[0]
        self.assertEqual(row['match_status'], 'missing')
        payload = [{'name': 'USD 10', 'count': 2}, {'name': 'USD 10.00', 'count': 8}]
        self.assertEqual(match_stock_targets([target], payload)[0]['match_status'], 'ambiguous')
        duplicate = {**target, 'nominal_id': 22}
        rows = match_stock_targets([target, duplicate], payload[:1])
        self.assertTrue(all(row['match_status'] == 'ambiguous' and row['stock_count'] is None for row in rows))

    def test_target_selection_skips_inactive_and_duplicate_nominals(self):
        # Повтор каталога не увеличивает число платёжных расчётов и запросов detail.
        services = catalog()
        disabled = {**services[0], 'raw': {'active': False}}
        self.assertEqual(collect_price_targets([disabled]), [])
        self.assertEqual(len(collect_price_targets(services + services)), 3)

    def test_cache_saves_error_after_success_with_new_date_and_null_count(self):
        # Проверяем SQL обновления: старая цифра и старая дата не сохраняются после неудачной проверки.
        # Контекстные менеджеры моделируем явно, чтобы не подключаться ни к одной БД.
        db = MagicMock()
        target = collect_price_targets(catalog())[:1]
        at = datetime(2026, 9, 10, 12, 0, tzinfo=timezone.utc)
        row = match_stock_targets(target, {'success': False}, checked_at=at)[0]
        save_stock_rows(db, 'unused', [row], 'batch', 'owner')
        conn = db.connect.return_value.__enter__.return_value
        sql, params = conn.cursor.return_value.__enter__.return_value.execute.call_args.args
        self.assertIn('stock_count=EXCLUDED.stock_count', sql)
        self.assertIn('checked_at=EXCLUDED.checked_at', sql)
        self.assertIsNone(params[4])
        self.assertEqual(params[5], 'error')
        self.assertEqual(params[9], at)
        conn.commit.assert_called_once()

    def test_export_keeps_zero_unknown_date_and_raw_error(self):
        # Лист остатков полезен даже при отсутствии цены; формулы из имени поставщика остаются текстом.
        targets = collect_price_targets(catalog())
        rows = match_stock_targets(targets[:2], [{'name': 'USD 10', 'count': 0}],
                                   checked_at=datetime(2026, 9, 10, 12, 0, tzinfo=timezone.utc))
        rows += match_stock_targets(targets[2:], {'success': False, 'message': 'Error'})
        rows[0]['service_title'] = '=1+1'
        workbook = load_workbook(BytesIO(build_interhub_prices_xlsx([], [], rows)))
        self.assertEqual(workbook.sheetnames, ['Закупочные цены', 'Ошибки calculate', 'Остатки'])
        sheet = workbook['Остатки']
        self.assertEqual(sheet['E2'].value, 0)
        self.assertIsNone(sheet['E3'].value)
        self.assertIsNone(sheet['E4'].value)
        self.assertEqual(sheet['F2'].value, '10.09.2026 15:00:00')
        self.assertIn('false', sheet['J4'].value)
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertEqual(sheet.freeze_panes, 'A2')

    def test_detail_get_uses_token_proxy_and_preserves_raw_error(self):
        # Новый метод пользуется тем же proxy и токеном; ответ с success=false передаётся в кэш целиком.
        service = build_interhub_service(HTTPException=HTTPException, interhub_api_url='https://interhub.test',
                                        interhub_token='test-token', timeout_sec=20, ssl_verify=True,
                                        ca_cert_path='', proxy_url='http://localhost:3128',
                                        calculate_path='/calculate', check_path='/check', deposit_path='/deposit')
        payload = {'success': False, 'status': -1, 'message': 'Error'}
        opener = MagicMock()
        opener.open.return_value.__enter__.return_value.read.return_value = json.dumps(payload).encode()
        with patch('domains.interhub_service.urllib.request.build_opener', return_value=opener):
            self.assertEqual(service.get_service_detail(9931), payload)
        request = opener.open.call_args.args[0]
        self.assertEqual(request.full_url, 'https://interhub.test/api/agent/service/detail?id=9931')
        self.assertEqual(request.get_method(), 'GET')
        self.assertEqual(request.get_header('Token'), 'test-token')
        self.assertIsNone(request.data)


class StockRefreshApiTests(unittest.TestCase):
    def test_payment_check_fetches_fresh_stock_each_time_without_paying(self):
        # Повторная проверка получает новый остаток по услуге и сохраняет обычный check для pay.
        detail = Mock(side_effect=[[{'name': 'USD 10', 'count': 5}], [{'name': 'USD 10', 'count': 0}]])
        client, db = self.make_client(detail=detail)
        self.check.return_value = {'success': True, 'message': 'Ready'}
        for count in (5, 0):
            response = client.post('/integrations/interhub/check', json={
                'service_id': 7, 'account': '', 'agent_transaction_id': f'check-{count}', 'params': {'nominal': 11},
            })
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.json()['success'])
            self.assertEqual(response.json()['stock']['stock_count'], count)
            self.assertEqual(response.json()['stock']['nominal_id'], 11)
            self.assertTrue(response.json()['stock']['checked_at'])
        self.assertEqual([call.args for call in detail.call_args_list], [(7,), (7,)])
        self.assertEqual(self.check.call_count, 2)
        self.pay.assert_not_called()
        sql = ' '.join(call.args[0] for call in db.connect.return_value.__enter__.return_value.cursor.return_value.__enter__.return_value.execute.call_args_list)
        self.assertIn('INSERT INTO app.interhub_transactions', sql)
        self.assertNotIn('interhub_stock_cache', sql)

    def test_payment_check_stock_errors_do_not_change_provider_availability(self):
        # Сетевой сбой, пустой ответ и success=false не превращаются в ноль и не меняют ответ check.
        cases = [Mock(side_effect=HTTPException(504, 'Timeout')),
                 Mock(return_value=[]), Mock(return_value={'success': False, 'message': 'No stock data'})]
        for detail in cases:
            for success in (True, False):
                client, _ = self.make_client(detail=detail)
                self.check.return_value = {'success': success, 'message': 'Check result'}
                response = client.post('/integrations/interhub/check', json={
                    'service_id': 7, 'account': '', 'agent_transaction_id': 'test-check', 'params': {'nominal': 11},
                })
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json()['success'], success)
                self.assertEqual(response.json()['message'], 'Check result')
                self.assertIsNone(response.json()['stock']['stock_count'])
                self.assertTrue(response.json()['stock']['message'])
                self.pay.assert_not_called()

    def test_payment_check_without_nominal_does_not_request_stock(self):
        # Обычное пополнение без номинала не запускает неподходящий метод остатков.
        client, _ = self.make_client()
        self.check.return_value = {'success': True}
        response = client.post('/integrations/interhub/check', json={
            'service_id': 7, 'account': 'test', 'agent_transaction_id': 'top-up', 'params': {},
        })
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(response.json()['stock'])
        self.detail.assert_not_called()
        self.get_services.assert_not_called()

    def make_client(self, role='owner', detail=None, calculate=None):
        # Собираем реальный обработчик с изолированными внешними вызовами и пустой базой цен.
        app = FastAPI()
        db = MagicMock()
        cur = db.connect.return_value.__enter__.return_value.cursor.return_value.__enter__.return_value
        cur.fetchall.return_value = []
        cur.fetchone.return_value = None
        user = SimpleNamespace(username=role)

        def require_role(*roles):
            # Проверяем права до запуска потока или чтения кэша для выгрузки.
            def authorize():
                # Только разрешённая роль получает доступ к операции.
                if role not in roles:
                    raise HTTPException(403)
                return user
            return authorize

        self.get_services = Mock(return_value=catalog())
        self.detail = detail or Mock(return_value=[{'name': 'USD 10', 'count': 0}, {'name': 'USD 20', 'count': 5}])
        self.calculate = calculate or Mock(return_value={'success': True, 'fixed_amount': 10})
        self.check, self.pay = Mock(), Mock()
        from domains.supplier_catalog import SupplierCatalog
        from domains.voucher_catalog_service import InterhubVoucherProvider
        cur.fetchone.side_effect = lambda: (True,) if cur.execute.call_args.args[0].startswith('SELECT pg_try_advisory') else None
        self.shared = SupplierCatalog(db, 'unused', InterhubVoucherProvider(self.get_services, self.detail, self.calculate, 0))
        self.shared.rows = Mock(return_value=[])
        mount_interhub_routes(app, DB_DSN='unused', psycopg=db, get_current_user=lambda: user,
                              require_role=require_role, UserOut=ApiModel, InterHubServiceListOut=ApiModel,
                              InterHubBalanceOut=ApiModel, InterHubPaymentRequestIn=ApiModel,
                              InterHubPaymentCheckOut=ApiModel, InterHubPayRequestIn=ApiModel,
                              InterHubVoucherBatchPayRequestIn=ApiModel, interhub_get_services=self.get_services,
                              interhub_get_balance=Mock(), interhub_calculate=self.calculate,
                              interhub_check=self.check, interhub_pay=self.pay, interhub_check_status=Mock(),
                              interhub_get_service_detail=self.detail, price_calculate_delay_ms=0, shared_catalog=self.shared)
        return TestClient(app), db

    def run_job(self, client):
        # Захватываем целевой worker и выполняем его синхронно, проверяя настоящий HTTP-прогресс.
        with patch('domains.interhub_api.threading.Thread') as thread:
            response = client.post('/integrations/interhub/prices/refresh')
            self.assertEqual(response.status_code, 200)
            job = response.json()
            self.assertEqual(client.post('/integrations/interhub/prices/refresh').status_code, 409)
            args = thread.call_args.kwargs
            args['target'](*args['args'])
        return client.get(f"/integrations/interhub/prices/refresh/{job['job_id']}").json()

    def test_refresh_one_detail_per_service_even_when_calculate_fails(self):
        # Пополнение пропускается в обоих этапах, ошибка цены не отменяет опрос двух ваучерных услуг.
        client, db = self.make_client(calculate=Mock(side_effect=[RuntimeError('calculate failed'), {'success': True, 'fixed_amount': 10}, {'success': True, 'fixed_amount': 10}]))
        self.get_services.return_value = catalog() + [{**catalog()[1], 'service_id': 9, 'type': 'VOUCHER'}]
        with patch.object(self.shared, 'write_snapshot', wraps=self.shared.write_snapshot) as save:
            job = self.run_job(client)
        self.assertEqual(job['state'], 'completed')
        self.assertEqual(job['processed'], 3)
        self.assertEqual(job['errors'], 1)
        self.assertEqual(job['stock_processed'], 2)
        self.assertEqual(job['stock_total'], 2)
        self.assertEqual(job['stock_successes'], 2)
        self.assertEqual(job['stock_errors'], 1)
        self.assertEqual([call.args[0] for call in self.detail.call_args_list], [7, 9])
        self.assertEqual([call.args[0]['service_id'] for call in self.calculate.call_args_list], [7, 7, 9])
        self.assertEqual(self.calculate.call_count, 3)
        self.assertEqual(save.call_count, 6)
        self.assertEqual([call.args[1]['service_id'] for call in save.call_args_list if call.args[2] == 'stocks'], [7, 7, 9])
        self.assertEqual(save.call_args_list[0].args[3]['value'], 0)
        self.check.assert_not_called()
        self.pay.assert_not_called()

    def test_http_200_provider_error_and_network_error_do_not_break_prices(self):
        # Оба вида ошибок сохраняются как неизвестный остаток, но все три calculate завершаются.
        client, _ = self.make_client(detail=Mock(side_effect=[{'success': False, 'message': 'Error'}, TimeoutError('timeout')]))
        self.get_services.return_value = catalog() + [{**catalog()[1], 'service_id': 9, 'type': 'VOUCHER'}]
        with patch.object(self.shared, 'write_snapshot', wraps=self.shared.write_snapshot) as save:
            job = self.run_job(client)
        self.assertEqual(job['successes'], 3)
        self.assertEqual(job['stock_errors'], 3)
        for call in save.call_args_list:
            if call.args[2] == 'stocks':
                self.assertTrue(call.args[3]['error'])

    def test_refresh_without_active_vouchers_skips_prices_and_stocks(self):
        # Пополнения и выключенные ваучеры не вызывают внешних запросов или записей в кэш.
        client, db = self.make_client()
        self.get_services.return_value = [
            catalog()[1], {**catalog()[0], 'active': False},
            {**catalog()[1], 'service_id': 9, 'type': 'TOP_UP'},
            {**catalog()[1], 'service_id': 10, 'type': ''},
            {**catalog()[1], 'service_id': 11, 'type': 'VOUCHER', 'fields': [
                {'name': 'nominal', 'value_list': [{'id': 1, 'title': 'Disabled', 'active': False}]}]},
        ]
        with patch.object(self.shared, 'write_snapshot', wraps=self.shared.write_snapshot) as save:
            job = self.run_job(client)
        self.assertEqual(job['state'], 'completed')
        self.assertEqual(job['total'], 0)
        self.assertEqual(job['processed'], 0)
        self.assertEqual(job['stock_total'], 0)
        self.assertEqual(job['stock_processed'], 0)
        self.detail.assert_not_called()
        self.calculate.assert_not_called()
        self.pay.assert_not_called()
        self.check.assert_not_called()
        save.assert_not_called()
        self.assertTrue(db.connect.called)

    def test_cached_json_and_export_do_not_contact_provider(self):
        # Повторное открытие вкладки и Excel читают один и тот же сохранённый ноль и дату.
        client, _ = self.make_client()
        stocks = match_stock_targets(collect_price_targets(catalog())[:1], [{'name': 'USD 10', 'count': 0}])
        with patch.object(self.shared, 'legacy', return_value=([], stocks)):
            latest = client.get('/integrations/interhub/prices/latest')
            exported = client.get('/integrations/interhub/prices/export')
        self.assertEqual(latest.status_code, 200)
        self.assertEqual(latest.json()['stocks'][0]['stock_count'], 0)
        self.assertTrue(latest.json()['stocks'][0]['checked_at'])
        self.assertEqual(exported.status_code, 200)
        self.assertEqual(load_workbook(BytesIO(exported.content))['Остатки']['E2'].value, 0)
        self.detail.assert_not_called()
        self.calculate.assert_not_called()
        self.get_services.assert_not_called()

    def test_operator_cannot_refresh_or_export(self):
        # Добавление остатков не расширяет права оператора на массовые запросы и выгрузку.
        client, _ = self.make_client(role='operator')
        with patch('domains.interhub_api.threading.Thread') as thread:
            self.assertEqual(client.post('/integrations/interhub/prices/refresh').status_code, 403)
            self.assertEqual(client.get('/integrations/interhub/prices/export').status_code, 403)
        thread.assert_not_called()


if __name__ == '__main__':
    unittest.main()
