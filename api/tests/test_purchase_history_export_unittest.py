import unittest
from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock

from fastapi import HTTPException
from openpyxl import load_workbook

from domains.purchase_history_export import build_purchase_history_xlsx


class PurchaseHistoryExportTests(unittest.TestCase):
    def test_exports_all_pages_of_both_sources_with_excel_types(self):
        # Общий лист должен точно объединять все успешные покупки двух источников без повторного чтения баз.
        crm_rows = [{"agent_transaction_id": f"crm-{i}", "service_id": 7, "nominal": "0015",
                     "price": 12.5, "created_at": "2026-09-04T21:00:00Z", "gift_code": "=1+1",
                     "deal_id": 42, "order_number": "000123"} for i in range(101)]
        hub_rows = [{"id": f"seller-{i}", "consumer_id": "seller", "service_id": 7,
                     "nominal_id": "0015", "amount": "14.25", "state": "succeeded",
                     "created_at": "2026-09-05T20:59:59Z", "gift_code": "SECRET"} for i in range(102)]
        hub_rows.extend({"id": f"seller-{state}", "service_id": 7, "state": state, "amount": "99"}
                        for state in ("failed", "processing", "requires_attention", "paid"))

        def hub_page(query):
            # Имитируем серверный фильтр до пагинации, чтобы не потерять успешные покупки на следующих страницах.
            filtered = [item for item in hub_rows if not query.get("state") or item["state"] == query["state"]]
            return {"total": len(filtered), "items": filtered[query["offset"]:query["offset"] + query["limit"]]}

        crm = Mock(return_value=iter(crm_rows))
        hub = SimpleNamespace(list_transactions=Mock(side_effect=hub_page), reveal_result=Mock())
        content = build_purchase_history_xlsx(
            load_crm_rows=crm, supplier_hub_client=hub,
            services=[{"service_id": 7, "title": "Steam", "fields": [
                {"name": "nominal", "value_list": [{"id": "0015", "title": "15 USD"}]}]}],
            date_from=date(2026, 9, 5), date_to=date(2026, 9, 5),
        )
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.sheetnames, ["CRM", "Селлер", "Все покупки"])
        self.assertEqual(workbook["CRM"].max_row, 102)
        self.assertEqual(workbook["Селлер"].max_row, 103)
        self.assertEqual(workbook["CRM"]["A102"].value, "crm-100")
        self.assertEqual(workbook["Селлер"]["A103"].value, "seller-101")
        self.assertEqual({cell.value for cell in workbook["Селлер"]["H"][1:]}, {"Выполнено"})
        self.assertEqual(workbook["CRM"]["I2"].value, datetime(2026, 9, 5))
        self.assertEqual(workbook["Селлер"]["I2"].value, datetime(2026, 9, 5, 23, 59, 59))
        self.assertEqual(workbook["CRM"]["G2"].value, 12.5)
        self.assertEqual(workbook["Селлер"]["G2"].value, 14.25)
        self.assertEqual(workbook["CRM"]["E2"].value, "0015")
        self.assertEqual(workbook["CRM"]["K2"].value, "000123")
        self.assertEqual(workbook["CRM"]["O2"].value, "=1+1")
        self.assertEqual(workbook["CRM"]["O2"].data_type, "s")
        self.assertIsNone(workbook["Селлер"]["O2"].value)
        self.assertEqual(workbook["Селлер"]["F2"].value, "15 USD")
        self.assertEqual(workbook["CRM"].freeze_panes, "A2")
        self.assertEqual(workbook["CRM"].auto_filter.ref, "A1:O102")
        combined = workbook["Все покупки"]
        separate_rows = [row for name in ("CRM", "Селлер") for row in workbook[name].iter_rows(min_row=2, values_only=True)]
        combined_rows = list(combined.iter_rows(min_row=2, values_only=True))
        self.assertEqual(combined_rows, separate_rows)
        self.assertEqual(combined.max_row, 204)
        self.assertEqual(sum(row[6] for row in combined_rows), 101 * 12.5 + 102 * 14.25)
        self.assertEqual({row[1] for row in combined_rows}, {"CRM", "seller"})
        self.assertEqual(combined["O2"].data_type, "s")
        self.assertEqual(combined["G2"].number_format, '#,##0.00')
        self.assertEqual(combined["I2"].number_format, 'dd.mm.yyyy hh:mm:ss')
        self.assertEqual(combined.freeze_panes, "A2")
        self.assertEqual(combined.auto_filter.ref, "A1:O204")
        crm.assert_called_once_with()
        self.assertEqual(hub.list_transactions.call_count, 2)
        for call in hub.list_transactions.call_args_list:
            self.assertEqual(call.args[0]["state"], "succeeded")
            self.assertEqual(call.args[0]["created_from"], "2026-09-05T00:00:00+03:00")
            self.assertEqual(call.args[0]["created_to"], "2026-09-06T00:00:00+03:00")
        hub.reveal_result.assert_not_called()

    def test_empty_period_still_contains_all_sheet_headers(self):
        # Пустая выборка сохраняет три листа с одинаковыми колонками.
        empty = Mock(return_value={"total": 0, "items": []})
        workbook = load_workbook(BytesIO(build_purchase_history_xlsx(
            load_crm_rows=lambda: iter(()), supplier_hub_client=SimpleNamespace(list_transactions=empty),
            services=[], date_from=None, date_to=None,
        )))
        self.assertTrue(all(sheet.max_row == 1 for sheet in workbook))
        self.assertEqual(workbook.sheetnames, ["CRM", "Селлер", "Все покупки"])
        headers = [tuple(cell.value for cell in sheet[1]) for sheet in workbook]
        self.assertEqual(headers, [headers[0]] * 3)

    def test_combined_sheet_preserves_purchases_when_one_source_is_empty(self):
        # Отсутствие покупок в одной базе не должно скрывать строки другой базы на общем листе.
        for source in ("CRM", "Селлер"):
            with self.subTest(source=source):
                purchase = {"id": "purchase-1", "agent_transaction_id": "purchase-1", "service_id": 7,
                            "state": "paid" if source == "CRM" else "succeeded", "price": 12.5, "amount": "12.5"}
                crm = [purchase] if source == "CRM" else []
                seller = [purchase] if source == "Селлер" else []
                workbook = load_workbook(BytesIO(build_purchase_history_xlsx(
                    load_crm_rows=lambda: iter(crm),
                    supplier_hub_client=SimpleNamespace(list_transactions=Mock(return_value={"total": len(seller), "items": seller})),
                    services=[], date_from=None, date_to=None,
                )))
                self.assertEqual(list(workbook["Все покупки"].values), list(workbook[source].values))
                self.assertEqual(workbook["Все покупки"].max_row, 2)

    def test_missing_page_or_unavailable_seller_does_not_return_partial_file(self):
        # Не выдаём успешный файл, если селлер оборвал историю или не ответил.
        for response in ({"total": 1, "items": []}, HTTPException(503, "unavailable")):
            hub = Mock(side_effect=response) if isinstance(response, Exception) else Mock(return_value=response)
            with self.subTest(response=response), self.assertRaises(HTTPException):
                build_purchase_history_xlsx(
                    load_crm_rows=Mock(return_value=iter(())),
                    supplier_hub_client=SimpleNamespace(list_transactions=hub),
                    services=[], date_from=None, date_to=None,
                )
