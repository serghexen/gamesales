from datetime import date,datetime,timezone
from decimal import Decimal
from io import BytesIO
import unittest
from unittest.mock import MagicMock,patch
from fastapi import FastAPI,HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from domains.voucher_sales_report import assemble_rows,build_xlsx,report_window,mount_voucher_sales_report,normalized_hash,read_seller,match_purchases


def sale(key=1,source='supplier',**extra):
    # Строка контракта содержит только метаданные заказа и отпечатки, без кода.
    return dict(key_id=key,connection_id=1,shop='JoyCards',order_id='00123',title='Apple 20 USD',
        delivered_at='2026-09-30T20:59:59+00:00',service_id=7,nominal_id='20',provider_code='interhub',
        amount='123.45',attempt_id=1,delivery_source=source,**extra)


class ReportTests(unittest.TestCase):
    def setUp(self):
        # Один и тот же календарный период используется всеми проверками сверки и Excel.
        self.start,self.end=report_window(date(2026,9,1),date(2026,9,30))
        self.catalog={('7','20'):('po_Apple','USD 20')}

    def assemble(self,rows,matches=None):
        # Проверяем чистую бизнес-логику независимо от HTTP и БД.
        return assemble_rows(rows,matches or {},self.catalog,{'7':'VOUCHER'},self.start,self.end)

    def purchase(self,hashed,amount='20.00'):
        # Покупка была сделана до выдачи и покрывает один код комплекта.
        return dict(id=hashed,amount=amount,hashes={hashed},created_at=datetime(2026,8,1,tzinfo=timezone.utc))

    def test_complete_bundle_partial_and_shared_purchase(self):
        # Только все уникальные коды дают полную цену; совпадение одной части не завышает итог.
        row=sale(source='manual',fingerprints={'whole':'whole','normalized':'normalized','tokens':['a','b']})
        matches={1:{'tx1':self.purchase('a'),'tx2':self.purchase('b')}}
        result=self.assemble([row],matches)[0]
        self.assertEqual(result['amount'],Decimal('40'))
        self.assertEqual(result['kind'],'manual')
        partial=self.assemble([row],{1:{'tx1':self.purchase('a')}})[0]
        self.assertIsNone(partial['amount']);self.assertIn('Найденная часть',partial['detail'])
        other={**row,'key_id':2}
        reused=self.assemble([row,other],{1:matches[1],2:matches[1]})
        self.assertTrue(all(r['amount'] is None for r in reused))

    def test_missing_and_future_purchase_not_confirmed(self):
        # Позднейшая закупка не может объяснять уже выданный код.
        row=sale(source='pool',fingerprints={'whole':'a','normalized':'a','tokens':['a']})
        future={**self.purchase('a'),'created_at':datetime(2026,10,1,tzinfo=timezone.utc)}
        self.assertIsNone(self.assemble([row],{1:{'tx':future}})[0]['amount'])
        self.assertEqual(self.assemble([row],{1:{'tx':self.purchase('a')}})[0]['kind'],'pool')

    def test_date_scope_duplicates_and_topups(self):
        # Верхняя граница исключительна; повторы отклоняются, пополнения не становятся ваучерами.
        with self.assertRaises(ValueError):self.assemble([sale(),sale()])
        with self.assertRaises(ValueError):self.assemble([{**sale(),'delivered_at':self.end.isoformat()}])
        self.assertEqual(assemble_rows([sale()],{},self.catalog,{'7':'TOP_UP'},self.start,self.end),[])

    def test_excel_layout_types_colors_names_and_formula_safety(self):
        # Имя магазина сразу перед таблицей; даты и деньги типизированы, формулы из данных отключены.
        rows=self.assemble([sale()]);rows[0].update(shop='=Bad/Shop',order='=1+1')
        missing={**rows[0],'connection_id':2,'kind':'unconfirmed','amount':None,'comment':'Не подтверждена','detail':'Проверить'}
        wb=load_workbook(BytesIO(build_xlsx(rows+[missing])))
        self.assertEqual(len(wb.worksheets),2);self.assertNotEqual(*wb.sheetnames)
        s=wb.worksheets[0]
        self.assertEqual(s['A1'].value,'=Bad/Shop');self.assertEqual(s['A1'].data_type,'s')
        self.assertEqual(s['A2'].value,'Номер заказа');self.assertEqual(s['A3'].data_type,'s')
        self.assertEqual(s['D3'].value,123.45);self.assertIsInstance(s['E3'].value,datetime)
        self.assertEqual(s['E3'].value.hour,23);self.assertEqual(s.freeze_panes,'A3')
        unknown=wb.worksheets[1];self.assertEqual(unknown['D4'].value,'Не подтверждена')
        self.assertTrue(unknown['A4'].fill.fgColor.rgb.endswith('FBE4DF'));self.assertEqual(unknown['F4'].comment.text,'Проверить')
        with self.assertRaises(HTTPException) as error:build_xlsx([])
        self.assertEqual(error.exception.status_code,404)

    def test_confirmed_table_and_separate_problem_block(self):
        # Неподтверждённая цена идёт вниз, за пределы фильтра основной таблицы; суммы и цвета сохраняются.
        normal=self.assemble([sale()])[0]
        missing={**normal,'order':'missing','amount':None,'kind':'unconfirmed',
                 'comment':'Ручная выдача — закупка не найдена.','detail':'Проверить закупку'}
        manual={**normal,'order':'manual','kind':'manual','comment':'Ручная выдача','detail':'Точный код'}
        pool={**normal,'order':'pool','kind':'pool','comment':'Из пула','detail':'Точный код'}
        sheet=load_workbook(BytesIO(build_xlsx([missing,normal,manual,pool]))).active
        table=list(sheet.tables.values())[0]
        self.assertEqual(table.ref,'A2:F5');self.assertEqual(table.tableStyleInfo.name,'TableStyleMedium2')
        self.assertTrue(table.tableStyleInfo.showRowStripes)
        self.assertEqual([sheet.cell(n,1).value for n in (3,4,5,10)],['00123','manual','pool','missing'])
        self.assertIsNone(sheet['A6'].value);self.assertIsNone(sheet['A7'].value)
        self.assertEqual(sheet['A8'].value,'Ручная выдача — цена не подтверждена')
        self.assertEqual(sheet['A9'].value,'Номер заказа');self.assertEqual(sheet['D10'].value,'Не подтверждена')
        self.assertEqual(sheet['F10'].comment.text,'Проверить закупку')
        for cell,color in [('A4','FFF1C9'),('A5','ECE3F8'),('A10','FBE4DF')]:
            self.assertTrue(sheet[cell].fill.fgColor.rgb.endswith(color))
        self.assertEqual(sheet['E3'].number_format,'dd/mm/yyyy hh:mm')
        self.assertEqual(sheet.row_dimensions[3].height,25)
        self.assertEqual(sheet['A2'].alignment.horizontal,'center')
        only=load_workbook(BytesIO(build_xlsx([missing]))).active
        self.assertEqual(len(only.tables),0);self.assertEqual(only.max_row,4)
        self.assertEqual(only['A2'].value,'Ручная выдача — цена не подтверждена')

    def test_route_owner_required_and_busy_retry(self):
        # Повторный одновременный запрос не запускает второе дорогое чтение архива.
        db=MagicMock();cur=db.connect.return_value.__enter__.return_value.cursor.return_value.__enter__.return_value
        cur.fetchone.return_value=(False,)
        def require_role(role):
            # Тестовая зависимость подтверждает назначенное право.
            self.assertEqual(role,'owner')
            return lambda: 'owner'
        app=FastAPI();mount_voucher_sales_report(app,db=db,dsn='test',require_role=require_role)
        with patch('domains.voucher_sales_report.read_seller') as read:
            response=TestClient(app).get('/integrations/interhub/voucher-sales/export?date_from=2026-09-01&date_to=2026-09-30')
            self.assertEqual(response.status_code,409);read.assert_not_called()

    def test_streamed_archive_matches_only_exact_codes(self):
        # Архив читается серверным курсором; совпадение номинала без кода не связывает закупку.
        cur=MagicMock();archive=cur.connection.cursor.return_value.__enter__.return_value
        hashed=normalized_hash('CODE-1234567890')
        row=sale(source='manual',fingerprints={'whole':hashed,'normalized':hashed,'tokens':[hashed]})
        stamp=datetime(2026,8,1,tzinfo=timezone.utc)
        archive.fetchmany.side_effect=[[
            ('yes',7,'20',Decimal('20'),stamp,'code1234567890'),
            ('no',7,'20',Decimal('20'),stamp,'other1234567890')],[]]
        cur.fetchall.return_value=[]
        matches=match_purchases(cur,[row],self.end)
        self.assertEqual(set(matches[1]),{'interhub:yes'})
        cur.connection.cursor.assert_called_once_with(name='voucher_report_archive')
        self.assertNotIn('code1234567890',str(matches))

    def test_download_and_source_failure(self):
        # Полный HTTP-путь создаёт открываемый Excel, ошибка источника не возвращает частичный файл.
        db=MagicMock();cur=db.connect.return_value.__enter__.return_value.cursor.return_value.__enter__.return_value
        cur.fetchone.return_value=(True,)
        app=FastAPI();mount_voucher_sales_report(app,db=db,dsn='test',require_role=lambda role: lambda: 'owner')
        client=TestClient(app);url='/integrations/interhub/voucher-sales/export?date_from=2026-09-01&date_to=2026-09-30'
        payload={'cutoff':self.end.isoformat(),'rows':[sale()]}
        with patch('domains.voucher_sales_report.read_seller',return_value=payload) as read, patch('domains.voucher_sales_report.load_catalog',return_value=(self.catalog,{'7':'VOUCHER'})), patch('domains.voucher_sales_report.match_purchases',return_value={}):
            response=client.get(url)
            self.assertEqual(response.status_code,200)
            self.assertEqual(response.headers['cache-control'],'no-store')
            self.assertIn('attachment',response.headers['content-disposition'])
            self.assertEqual(load_workbook(BytesIO(response.content)).active['D3'].value,123.45)
            read.side_effect=RuntimeError('private source failure')
            response=client.get(url)
            self.assertEqual(response.status_code,502)
            self.assertNotIn('private',response.text)

    def test_https_api_prefix(self):
        # Production nginx публикует Seller под /api; секрет отправляется только на этот проверенный путь.
        payload={'version':1,'rows':[],'total':0,'date_from':'2026-09-01','date_to':'2026-09-30'}
        import json
        opener=MagicMock();opener.open.return_value.__enter__.return_value.read.return_value=json.dumps(payload).encode()
        with patch.dict('os.environ',{'SELLER_REPORT_URL':'https://seller.example.com/api/','SELLER_CRM_REPORT_TOKEN':'a'*32}), patch('domains.voucher_sales_report.urllib.request.build_opener',return_value=opener):
            self.assertEqual(read_seller(date(2026,9,1),date(2026,9,30)),payload)
        self.assertEqual(opener.open.call_args.args[0].full_url,'https://seller.example.com/api/internal/reports/voucher-sales?date_from=2026-09-01&date_to=2026-09-30')

    def test_unsafe_url_and_missing_configuration(self):
        # Секрет не уходит на HTTP внешнего сервера и в URL с чужими реквизитами.
        for url in ['http://example.org','https://user:password@example.org','https://example.org/path']:
            with patch.dict('os.environ',{'SELLER_REPORT_URL':url,'SELLER_CRM_REPORT_TOKEN':'a'*32}):
                with self.assertRaises(HTTPException) as error:read_seller(date(2026,9,1),date(2026,9,30))
                self.assertEqual(error.exception.status_code,503)


if __name__=='__main__':unittest.main()
